from datetime import datetime, time
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.db.models import DecimalField, IntegerField, Q, Sum
from django.db.models.functions import Coalesce
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from catalog.models import Product
from orders.models import Order, OrderItem


SALE_STATUSES = (Order.Status.PAID, Order.Status.DELIVERED)


def _bounds(date_from, date_to):
    start = timezone.make_aware(datetime.combine(date_from, time.min))
    end = timezone.make_aware(datetime.combine(date_to, time.max))
    return start, end


def build_report(date_from, date_to, stock_date):
    start, end = _bounds(date_from, date_to)
    stock_end = timezone.make_aware(datetime.combine(stock_date, time.max))
    sold_items = list(
        OrderItem.objects.filter(order__status__in=SALE_STATUSES, order__paid_at__range=(start, end))
        .select_related("order", "order__user", "product")
        .order_by("order__paid_at", "order_id")
    )
    product_metrics = Product.objects.filter(is_active=True).annotate(
        sold_quantity=Coalesce(
            Sum("order_items__quantity", filter=Q(order_items__order__status__in=SALE_STATUSES, order_items__order__paid_at__range=(start, end))),
            0,
            output_field=IntegerField(),
        ),
        sales_revenue=Coalesce(
            Sum("order_items__subtotal", filter=Q(order_items__order__status__in=SALE_STATUSES, order_items__order__paid_at__range=(start, end))),
            Decimal("0.00"),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        ),
    )
    top_products = list(product_metrics.order_by("-sold_quantity", "-sales_revenue", "name")[:10])
    unpopular_products = list(product_metrics.order_by("sold_quantity", "sales_revenue", "name")[:10])
    stock_rows = list(
        Product.objects.select_related("category", "manufacturer")
        .annotate(
            as_of_stock=Coalesce(
                Sum("stock_movements__delta", filter=Q(stock_movements__created_at__lte=stock_end)),
                0,
                output_field=IntegerField(),
            )
        )
        .order_by("name")
    )
    return {
        "date_from": date_from,
        "date_to": date_to,
        "stock_date": stock_date,
        "sold_items": sold_items,
        "sold_quantity": sum(item.quantity for item in sold_items),
        "revenue": sum((item.subtotal for item in sold_items), Decimal("0.00")),
        "orders_count": len({item.order_id for item in sold_items}),
        "stock_rows": stock_rows,
        "top_products": top_products,
        "unpopular_products": unpopular_products,
    }


def report_to_xlsx(data):
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="324F5E")
    header_font = Font(color="FFFFFF", bold=True)

    def add_sheet(title, headers, rows, widths=None):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate(widths or [18] * len(headers), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        return sheet

    sales = add_sheet(
        "Продажи",
        ["Заказ", "Дата оплаты", "Покупатель", "Артикул", "Товар", "Количество", "Цена, ₽", "Сумма, ₽"],
        [[item.order.number, timezone.localtime(item.order.paid_at).strftime("%d.%m.%Y %H:%M"), item.order.user.get_full_name() or item.order.user.username, item.sku, item.product_name, item.quantity, float(item.unit_price), float(item.subtotal)] for item in data["sold_items"]],
        [14, 20, 24, 18, 42, 14, 14, 16],
    )
    for row in sales.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = '#,##0.00 "₽"'
    add_sheet(
        "Остатки",
        ["Артикул", "Товар", "Категория", "Производитель", f"Остаток на {data['stock_date'].strftime('%d.%m.%Y')}"],
        [[product.sku, product.name, product.category.name, product.manufacturer.name, product.as_of_stock] for product in data["stock_rows"]],
        [18, 42, 24, 24, 20],
    )
    add_sheet("Популярные", ["Место", "Артикул", "Товар", "Продано", "Выручка, ₽"], [[index, product.sku, product.name, product.sold_quantity, float(product.sales_revenue)] for index, product in enumerate(data["top_products"], 1)], [10, 18, 42, 14, 18])
    add_sheet("Невостребованные", ["Место", "Артикул", "Товар", "Продано", "Выручка, ₽"], [[index, product.sku, product.name, product.sold_quantity, float(product.sales_revenue)] for index, product in enumerate(data["unpopular_products"], 1)], [10, 18, 42, 14, 18])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def report_to_pdf(data):
    font_dir = Path(settings.BASE_DIR) / "static" / "fonts"
    pdfmetrics.registerFont(TTFont("PTSans", str(font_dir / "PTSans-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("PTSans-Bold", str(font_dir / "PTSans-Bold.ttf")))
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title="Отчёт AutoParts Hub")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("RuTitle", parent=styles["Title"], fontName="PTSans-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#213943"), alignment=TA_CENTER)
    heading_style = ParagraphStyle("RuHeading", parent=styles["Heading2"], fontName="PTSans-Bold", fontSize=13, leading=16, textColor=colors.HexColor("#213943"), spaceAfter=7)
    body_style = ParagraphStyle("RuBody", parent=styles["BodyText"], fontName="PTSans", fontSize=8, leading=10)

    def paragraph(value):
        return Paragraph(str(value), body_style)

    def add_table(story, heading, headers, rows, widths):
        story.extend((Paragraph(heading, heading_style),))
        table = Table([[paragraph(value) for value in headers]] + [[paragraph(value) for value in row] for row in rows], colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#324F5E")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "PTSans-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "PTSans"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D4C8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F4ED")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend((table, Spacer(1, 6 * mm)))

    story = [
        Paragraph("AutoParts Hub — сводный отчёт", title_style),
        Paragraph(f"Период продаж: {data['date_from'].strftime('%d.%m.%Y')}–{data['date_to'].strftime('%d.%m.%Y')} · Остатки на {data['stock_date'].strftime('%d.%m.%Y')}", body_style),
        Spacer(1, 5 * mm),
    ]
    add_table(story, f"Проданные товары — {data['sold_quantity']} шт., {data['revenue']:.2f} ₽", ["Заказ", "Оплата", "Артикул", "Товар", "Кол-во", "Цена", "Сумма"], [[item.order.number, timezone.localtime(item.order.paid_at).strftime("%d.%m.%Y"), item.sku, item.product_name, item.quantity, f"{item.unit_price:.2f} ₽", f"{item.subtotal:.2f} ₽"] for item in data["sold_items"]] or [["—", "—", "—", "За период продаж нет", "—", "—", "—"]], [24 * mm, 25 * mm, 29 * mm, 78 * mm, 19 * mm, 26 * mm, 28 * mm])
    story.append(PageBreak())
    add_table(story, "Остатки товаров", ["Артикул", "Товар", "Категория", "Производитель", "Остаток"], [[product.sku, product.name, product.category.name, product.manufacturer.name, product.as_of_stock] for product in data["stock_rows"]], [32 * mm, 86 * mm, 45 * mm, 45 * mm, 28 * mm])
    story.append(PageBreak())
    add_table(story, "Самые популярные товары", ["№", "Артикул", "Товар", "Продано", "Выручка"], [[index, product.sku, product.name, product.sold_quantity, f"{product.sales_revenue:.2f} ₽"] for index, product in enumerate(data["top_products"], 1)], [15 * mm, 35 * mm, 100 * mm, 30 * mm, 45 * mm])
    add_table(story, "Самые невостребованные товары", ["№", "Артикул", "Товар", "Продано", "Выручка"], [[index, product.sku, product.name, product.sold_quantity, f"{product.sales_revenue:.2f} ₽"] for index, product in enumerate(data["unpopular_products"], 1)], [15 * mm, 35 * mm, 100 * mm, 30 * mm, 45 * mm])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("PTSans", 7)
        canvas.setFillColor(colors.HexColor("#65747A"))
        canvas.drawString(12 * mm, 7 * mm, "AutoParts Hub · учебный проект")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Страница {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
