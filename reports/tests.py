from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Category, Manufacturer, Product, StockMovement
from orders.services import add_to_cart, create_order, process_payment

from .services import build_report, report_to_pdf, report_to_xlsx


class ReportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username="admin", email="admin@example.com", password="pass12345", is_staff=True)
        self.buyer = User.objects.create_user(username="buyer", email="buyer@example.com", password="pass12345")
        category = Category.objects.create(name="Фильтры")
        manufacturer = Manufacturer.objects.create(name="Bosch", country="Германия")
        self.sold = Product.objects.create(name="Масляный фильтр", category=category, manufacturer=manufacturer, sku="F-1", price=Decimal("500.00"), stock=10, description="Описание")
        self.unpopular = Product.objects.create(name="Воздушный фильтр", category=category, manufacturer=manufacturer, sku="F-2", price=Decimal("700.00"), stock=4, description="Описание")
        add_to_cart(self.buyer, self.sold, 2)
        self.order = create_order(self.buyer, "Адрес доставки")
        process_payment(self.order, "4242424242424242")
        self.order.refresh_from_db()

    def test_report_totals_and_zero_sales_products(self):
        today = timezone.localdate()
        report = build_report(today - timedelta(days=1), today, today)
        self.assertEqual(report["sold_quantity"], 2)
        self.assertEqual(report["revenue"], Decimal("1000.00"))
        self.assertEqual(report["top_products"][0].pk, self.sold.pk)
        self.assertEqual(report["unpopular_products"][0].pk, self.unpopular.pk)
        self.assertEqual(report["unpopular_products"][0].sold_quantity, 0)

    def test_historical_stock_uses_movement_ledger(self):
        three_days_ago = timezone.now() - timedelta(days=3)
        yesterday = timezone.now() - timedelta(days=1)
        StockMovement.objects.filter(product=self.sold, reason=StockMovement.Reason.INITIAL).update(created_at=three_days_ago)
        StockMovement.objects.filter(order=self.order, product=self.sold).update(created_at=yesterday)
        report = build_report(timezone.localdate() - timedelta(days=5), timezone.localdate(), (timezone.localdate() - timedelta(days=2)))
        row = next(item for item in report["stock_rows"] if item.pk == self.sold.pk)
        self.assertEqual(row.as_of_stock, 10)

    def test_excel_and_pdf_exports_are_valid(self):
        today = timezone.localdate()
        report = build_report(today - timedelta(days=1), today, today)
        xlsx = report_to_xlsx(report)
        workbook = load_workbook(BytesIO(xlsx))
        self.assertEqual(workbook.sheetnames, ["Продажи", "Остатки", "Популярные", "Невостребованные"])
        pdf = report_to_pdf(report)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreater(len(pdf), 5000)

    def test_report_page_and_download_require_admin(self):
        self.client.force_login(self.buyer)
        self.assertEqual(self.client.get(reverse("staff:reports")).status_code, 403)
        self.client.force_login(self.admin)
        today = timezone.localdate().isoformat()
        response = self.client.get(reverse("staff:reports"), {"date_from": today, "date_to": today, "stock_date": today, "format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
